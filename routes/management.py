"""
Management portal routes for Moulya College Management System
Handles all management functionality
"""

from flask import Blueprint, render_template, request, redirect, url_for, flash, session, jsonify
from routes.auth import login_required
from services.management_service import ManagementService
from services.auth_service import AuthService
from services.reporting_service import ReportingService
from services.excel_export_service import ExcelExportService
from models.user import Lecturer
from models.academic import Course, Subject
from models.student import Student
from database import db
from utils.validators import validate_username, validate_password

def is_ajax_request():
    """Check if the current request is an AJAX request"""
    return (request.headers.get('X-Requested-With') == 'XMLHttpRequest' or 
            request.headers.get('Content-Type') == 'application/json' or
            request.is_json)

management_bp = Blueprint('management', __name__)

@management_bp.route('/dashboard')
@login_required('management')
def dashboard():
    """Management dashboard with overview statistics"""
    try:
        # Get overview statistics
        stats = ManagementService.get_dashboard_stats()
        return render_template('management/dashboard.html', stats=stats)
    except Exception as e:
        flash(f'Error loading dashboard: {str(e)}', 'error')
        return render_template('management/dashboard.html', stats={})

@management_bp.route('/lecturers')
@login_required('management')
def lecturers():
    """List all lecturers"""
    try:
        page = request.args.get('page', 1, type=int)
        search = request.args.get('search', '', type=str)
        
        lecturers_data = ManagementService.get_lecturers_paginated(page, search)
        subjects = Subject.query.filter_by(is_active=True).order_by(Subject.name).all()
        return render_template('management/lecturers.html', 
                             lecturers=lecturers_data['lecturers'],
                             pagination=lecturers_data['pagination'],
                             search=search,
                             subjects=subjects)
    except Exception as e:
        flash(f'Error loading lecturers: {str(e)}', 'error')
        subjects = Subject.query.filter_by(is_active=True).order_by(Subject.name).all()
        return render_template('management/lecturers.html', lecturers=[], pagination=None, subjects=subjects)

@management_bp.route('/students')
@login_required('management')
def students():
    """List all students"""
    try:
        page = request.args.get('page', 1, type=int)
        search = request.args.get('search', '', type=str)
        course_id = request.args.get('course_id', type=int)
        
        students_data = ManagementService.get_students_paginated(page, search, course_id)
        courses = Course.query.filter_by(is_active=True).all()
        
        return render_template('management/students.html',
                             students=students_data['students'],
                             pagination=students_data['pagination'],
                             courses=courses,
                             search=search,
                             selected_course=course_id)
    except Exception as e:
        flash(f'Error loading students: {str(e)}', 'error')
        return render_template('management/students.html', students=[], pagination=None, courses=[])

@management_bp.route('/courses')
@login_required('management')
def courses():
    """Manage courses"""
    try:
        page = request.args.get('page', 1, type=int)
        search = request.args.get('search', '', type=str)
        
        query = Course.query.filter_by(is_active=True)
        
        if search:
            query = query.filter(
                db.or_(
                    Course.name.contains(search),
                    Course.code.contains(search)
                )
            )
        
        pagination = query.order_by(Course.name).paginate(
            page=page, per_page=20, error_out=False
        )
        
        return render_template('management/courses.html', 
                             courses=pagination.items,
                             pagination=pagination,
                             search=search)
    except Exception as e:
        flash(f'Error loading courses: {str(e)}', 'error')
        return render_template('management/courses.html', courses=[], pagination=None, search='')

@management_bp.route('/subjects')
@login_required('management')
def subjects():
    """Manage subjects"""
    try:
        course_id = request.args.get('course_id', type=int)
        subjects_data = ManagementService.get_subjects_by_course(course_id)
        courses = Course.query.filter_by(is_active=True).all()
        
        return render_template('management/subjects.html',
                             subjects=subjects_data,
                             courses=courses,
                             selected_course=course_id)
    except Exception as e:
        flash(f'Error loading subjects: {str(e)}', 'error')
        return render_template('management/subjects.html', subjects=[], courses=[])



@management_bp.route('/lecturers/add', methods=['POST'])
@login_required('management')
def add_lecturer():
    """Add single lecturer"""
    try:
        lecturer_data = {
            'lecturer_id': request.form.get('lecturer_id', '').strip(),
            'name': request.form.get('name', '').strip(),
            'subject_ids': request.form.getlist('subject_ids')
        }
        
        success, message = ManagementService.add_lecturer(lecturer_data)
        
        # Handle AJAX requests
        if is_ajax_request():
            return jsonify({'success': success, 'message': message})
        
        if success:
            flash(message, 'success')
        else:
            flash(message, 'error')
            
    except Exception as e:
        error_msg = f'Error adding lecturer: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.lecturers'))

@management_bp.route('/lecturers/bulk', methods=['POST'])
@login_required('management')
def bulk_add_lecturers():
    """Bulk add lecturers from Excel file"""
    try:
        if 'file' not in request.files:
            error_msg = 'No file selected'
            if is_ajax_request():
                return jsonify({'success': False, 'message': error_msg})
            flash(error_msg, 'error')
            return redirect(url_for('management.lecturers'))

        file = request.files['file']
        if file.filename == '':
            error_msg = 'No file selected'
            if is_ajax_request():
                return jsonify({'success': False, 'message': error_msg})
            flash(error_msg, 'error')
            return redirect(url_for('management.lecturers'))

        if not file.filename.lower().endswith(('.xlsx', '.xls')):
            error_msg = 'Please upload an Excel file (.xlsx or .xls)'
            if is_ajax_request():
                return jsonify({'success': False, 'message': error_msg})
            flash(error_msg, 'error')
            return redirect(url_for('management.lecturers'))

        success, message, credentials, errors = ManagementService.bulk_add_lecturers(file.read())

        # Handle AJAX requests
        if is_ajax_request():
            return jsonify({'success': success, 'message': message, 'errors': errors})

        if success:
            flash(message, 'success')
            if errors:
                flash(f'Some errors occurred: {"; ".join(errors[:5])}', 'warning')
        else:
            flash(message, 'error')
            if errors:
                flash(f'Errors: {"; ".join(errors[:5])}', 'error')

    except Exception as e:
        error_msg = f'Error processing file: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')

    return redirect(url_for('management.lecturers'))

@management_bp.route('/lecturers/<int:lecturer_id>/toggle-status', methods=['POST'])
@login_required('management')
def toggle_lecturer_status(lecturer_id):
    """On deactivate: permanently delete the lecturer and dependencies; on inactive -> active: reactivate.

    This preserves the existing button/JS flow while ensuring a hard delete when switching from
    active to inactive.
    """
    try:
        lecturer = Lecturer.query.get_or_404(lecturer_id)

        if lecturer.is_active:
            # Deactivation request -> permanently delete
            success, message = ManagementService.delete_lecturer_permanently(lecturer_id)
            if not success:
                if is_ajax_request():
                    return jsonify({'success': False, 'message': message})
                flash(message, 'error')
                return redirect(url_for('management.lecturers'))

            if is_ajax_request():
                return jsonify({'success': True, 'message': message})
            flash(message, 'success')
            return redirect(url_for('management.lecturers'))
        else:
            # Inactive -> Active toggle retains record
            lecturer.is_active = True
            db.session.commit()
            message = f'Lecturer {lecturer.name} has been activated'
            if is_ajax_request():
                return jsonify({'success': True, 'message': message})
            flash(message, 'success')
            return redirect(url_for('management.lecturers'))

    except Exception as e:
        error_msg = f'Error updating lecturer status: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.lecturers'))

@management_bp.route('/lecturers/<int:lecturer_id>/reset-password', methods=['POST'])
@login_required('management')
def reset_lecturer_password(lecturer_id):
    """Reset lecturer password. If username/password provided, set custom values with validation."""
    try:
        lecturer = Lecturer.query.get_or_404(lecturer_id)

        payload = {}
        if request.is_json:
            payload = request.get_json(silent=True) or {}
        username = (payload.get('username') or request.form.get('username') or '').strip()
        password = (payload.get('password') or request.form.get('password') or '').strip()

        # If either field present, treat as custom update
        if username or password:
            # Validate fields
            if username:
                valid, msg = validate_username(username)
                if not valid:
                    if is_ajax_request():
                        return jsonify({'success': False, 'message': msg}), 400
                    flash(msg, 'error')
                    return redirect(url_for('management.lecturers'))
                # Uniqueness check excluding current lecturer
                existing = Lecturer.query.filter(Lecturer.username == username, Lecturer.id != lecturer.id).first()
                if existing:
                    msg = 'Username already exists. Please choose a different one.'
                    if is_ajax_request():
                        return jsonify({'success': False, 'message': msg}), 400
                    flash(msg, 'error')
                    return redirect(url_for('management.lecturers'))

            if password:
                valid, msg = validate_password(password)
                if not valid:
                    if is_ajax_request():
                        return jsonify({'success': False, 'message': msg}), 400
                    flash(msg, 'error')
                    return redirect(url_for('management.lecturers'))

            # Apply updates
            if username:
                lecturer.username = username
            if password:
                lecturer.set_password(password)
            db.session.commit()

            message = 'Credentials updated successfully'
            if is_ajax_request():
                return jsonify({'success': True, 'message': message})
            flash(message, 'success')
            return redirect(url_for('management.lecturers'))

        # Default behavior: auto-generate a new password
        success, new_password, message = AuthService.reset_lecturer_password(lecturer.lecturer_id)
        if success:
            response_message = f'Password reset for {lecturer.name}. New password: {new_password}'
            if is_ajax_request():
                return jsonify({'success': True, 'message': response_message, 'new_password': new_password})
            flash(response_message, 'success')
        else:
            if is_ajax_request():
                return jsonify({'success': False, 'message': message})
            flash(message, 'error')

    except Exception as e:
        error_msg = f'Error resetting password: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.lecturers'))

@management_bp.route('/lecturers/reset-passwords-all', methods=['POST'])
@login_required('management')
def reset_passwords_all_lecturers():
    """Set a custom password for ALL lecturers"""
    try:
        # Accept JSON or form data
        password = None
        if request.is_json:
            payload = request.get_json(silent=True) or {}
            password = (payload.get('password') or '').strip()
        if not password:
            password = (request.form.get('password') or '').strip()

        if not password:
            message = 'Password is required'
            if is_ajax_request():
                return jsonify({'success': False, 'message': message}), 400
            flash(message, 'error')
            return redirect(url_for('management.lecturers'))

        # Basic validation
        if len(password) < 6:
            message = 'Password must be at least 6 characters'
            if is_ajax_request():
                return jsonify({'success': False, 'message': message}), 400
            flash(message, 'error')
            return redirect(url_for('management.lecturers'))

        from models.user import Lecturer
        from database import db

        lecturers = Lecturer.query.all()
        for lecturer in lecturers:
            lecturer.set_password(password)
        db.session.commit()

        message = f'Password updated for {len(lecturers)} lecturer(s).'
        if is_ajax_request():
            return jsonify({'success': True, 'message': message})
        flash(message, 'success')
    except Exception as e:
        try:
            from database import db as _db
            _db.session.rollback()
        except Exception:
            pass
        error_msg = f'Error resetting passwords: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg}), 500
        flash(error_msg, 'error')

    return redirect(url_for('management.lecturers'))

@management_bp.route('/lecturers/credentials')
@login_required('management')
def view_bulk_credentials():
    """View bulk upload credentials"""
    credentials = session.pop('bulk_credentials', [])
    return render_template('management/bulk_credentials.html', credentials=credentials)

@management_bp.route('/lecturers/<int:lecturer_id>/password')
@login_required('management')
def get_lecturer_password(lecturer_id):
    """Get lecturer's decrypted password for management"""
    try:
        lecturer = Lecturer.query.get_or_404(lecturer_id)
        password = lecturer.get_decrypted_password()
        
        if password:
            return jsonify({'success': True, 'password': password})
        else:
            return jsonify({'success': False, 'message': 'Password not available'})
            
    except Exception as e:
        return jsonify({'success': False, 'message': f'Error retrieving password: {str(e)}'})

@management_bp.route('/students/<int:student_id>/password')
@login_required('management')
def get_student_password(student_id):
    """Get student's decrypted password for management"""
    try:
        student = Student.query.get_or_404(student_id)
        password = student.get_decrypted_password()
        
        if password:
            return jsonify({'success': True, 'password': password})
        else:
            return jsonify({'success': False, 'message': 'Password not available'})
            
    except Exception as e:
        return jsonify({'success': False, 'message': f'Error retrieving password: {str(e)}'})

@management_bp.route('/lecturers/<int:lecturer_id>/assign-subjects', methods=['POST'])
@login_required('management')
def assign_subjects_to_lecturer(lecturer_id):
    """Assign subjects to a lecturer"""
    try:
        subject_ids = request.form.getlist('subject_ids')
        
        if not subject_ids:
            if is_ajax_request():
                return jsonify({'success': False, 'message': 'Please select at least one subject'})
            flash('Please select at least one subject', 'error')
            return redirect(url_for('management.lecturers'))
        
        success, message = ManagementService.assign_subjects_to_lecturer(lecturer_id, subject_ids)
        
        if is_ajax_request():
            return jsonify({'success': success, 'message': message})
        
        if success:
            flash(message, 'success')
        else:
            flash(message, 'error')
        
        return redirect(url_for('management.lecturers'))
        
    except Exception as e:
        error_msg = f'Error assigning subjects: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
        return redirect(url_for('management.lecturers'))

@management_bp.route('/lecturers/<int:lecturer_id>/unassign-subject', methods=['POST'])
@login_required('management')
def unassign_subject_from_lecturer(lecturer_id):
    """Unassign a subject from a lecturer (current academic year)"""
    try:
        subject_id = request.form.get('subject_id')
        if not subject_id:
            if is_ajax_request():
                return jsonify({'success': False, 'message': 'Subject is required'})
            flash('Subject is required', 'error')
            return redirect(url_for('management.lecturers'))

        success, message = ManagementService.unassign_subject_from_lecturer(lecturer_id, subject_id)

        if is_ajax_request():
            return jsonify({'success': success, 'message': message})

        if success:
            flash(message, 'success')
        else:
            flash(message, 'error')

    except Exception as e:
        error_msg = f'Error unassigning subject: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')

    return redirect(url_for('management.lecturers'))

@management_bp.route('/lecturers/credentials/export')
@login_required('management')
def export_lecturer_credentials():
    """Export lecturer credentials to Excel"""
    try:
        from flask import make_response
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from io import BytesIO
        
        # Get all active lecturers with their credentials, sorted by lecturer_id then name
        lecturers = Lecturer.query.filter_by(is_active=True).order_by(Lecturer.lecturer_id.asc(), Lecturer.name.asc()).all()
        
        # Create workbook and worksheet
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Lecturer Credentials"
        
        # Headers
        headers = ['Lecturer ID', 'Name', 'Username', 'Password', 'Assigned Subjects', 'Created Date']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # Data rows
        for row, lecturer in enumerate(lecturers, 2):
            ws.cell(row=row, column=1, value=lecturer.lecturer_id)
            ws.cell(row=row, column=2, value=lecturer.name)
            ws.cell(row=row, column=3, value=lecturer.username)
            ws.cell(row=row, column=4, value=lecturer.get_decrypted_password() or 'N/A')
            assigned_subjects = sorted([subject.name for subject in lecturer.get_assigned_subjects()], key=lambda s: s.upper())
            ws.cell(row=row, column=5, value=', '.join(assigned_subjects) if assigned_subjects else 'No subjects assigned')
            ws.cell(row=row, column=6, value=lecturer.created_at.strftime('%Y-%m-%d') if lecturer.created_at else '')
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Create response
        response = make_response(output.getvalue())
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = 'attachment; filename=lecturer_credentials.xlsx'
        
        return response
        
    except Exception as e:
        flash(f'Error exporting credentials: {str(e)}', 'error')
        return redirect(url_for('management.lecturers'))

@management_bp.route('/students/credentials/export')
@login_required('management')
def export_student_credentials():
    """Export student credentials to Excel"""
    try:
        from flask import make_response
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from io import BytesIO
        
        course_id = request.args.get('course_id', type=int)
        
        # Get active students, optionally filtered by course, sorted by roll_number then name
        query = Student.query.filter_by(is_active=True)
        if course_id:
            query = query.filter_by(course_id=course_id)
        
        students = query.order_by(Student.roll_number.asc(), Student.name.asc()).all()
        
        # Create workbook and worksheet
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Student Credentials"
        
        # Headers - simplified as requested: Name, Username, Password
        headers = ['Name', 'Username', 'Password']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # Data rows
        for row, student in enumerate(students, 2):
            ws.cell(row=row, column=1, value=student.name)
            ws.cell(row=row, column=2, value=student.username)
            ws.cell(row=row, column=3, value=student.get_decrypted_password() or 'N/A')
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Create response
        response = make_response(output.getvalue())
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        
        # Set filename based on whether it's filtered by course
        if course_id:
            course = Course.query.get(course_id)
            filename = f'student_credentials_{course.code if course else "course"}.xlsx'
        else:
            filename = 'student_credentials.xlsx'
        
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        
        return response
        
    except Exception as e:
        flash(f'Error exporting credentials: {str(e)}', 'error')
        return redirect(url_for('management.students'))

@management_bp.route('/students/add', methods=['POST'])
@login_required('management')
def add_student():
    """Add single student"""
    try:
        student_data = {
            'roll_number': request.form.get('roll_number', '').strip(),
            'name': request.form.get('name', '').strip(),
            'course_id': int(request.form.get('course_id')),
            'academic_year': int(request.form.get('academic_year')),
            'current_semester': int(request.form.get('current_semester', 1)),
            'email': request.form.get('email', '').strip() or None,
            'date_of_birth': request.form.get('date_of_birth', '').strip() or None
        }
        
        success, message = ManagementService.add_student(student_data)
        
        # Handle AJAX requests
        if is_ajax_request():
            return jsonify({'success': success, 'message': message})
        
        if success:
            flash(message, 'success')
        else:
            flash(message, 'error')
            
    except Exception as e:
        error_msg = f'Error adding student: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.students'))

@management_bp.route('/students/bulk', methods=['POST'])
@login_required('management')
def bulk_add_students():
    """Bulk add students from Excel file"""
    try:
        if 'file' not in request.files:
            error_msg = 'No file selected'
            if is_ajax_request():
                return jsonify({'success': False, 'message': error_msg})
            flash(error_msg, 'error')
            return redirect(url_for('management.students'))
        
        file = request.files['file']
        if file.filename == '':
            error_msg = 'No file selected'
            if is_ajax_request():
                return jsonify({'success': False, 'message': error_msg})
            flash(error_msg, 'error')
            return redirect(url_for('management.students'))
        
        if not file.filename.lower().endswith(('.xlsx', '.xls')):
            error_msg = 'Please upload an Excel file (.xlsx or .xls)'
            if is_ajax_request():
                return jsonify({'success': False, 'message': error_msg})
            flash(error_msg, 'error')
            return redirect(url_for('management.students'))
        
        success, message, errors, credentials = ManagementService.bulk_add_students(file.read())
        
        # Handle AJAX requests
        if is_ajax_request():
            return jsonify({'success': success, 'message': message, 'errors': errors, 'credentials': credentials})
        
        if success:
            flash(message, 'success')
            if credentials:
                creds_msg = "Generated credentials: " + ", ".join([f"{c['roll_number']}:{c['password']}" for c in credentials[:5]])
                if len(credentials) > 5:
                    creds_msg += f" (+{len(credentials)-5} more)"
                flash(creds_msg, 'info')
            if errors:
                flash(f'Some errors occurred: {"; ".join(errors[:5])}', 'warning')
        else:
            flash(message, 'error')
            if errors:
                flash(f'Errors: {"; ".join(errors[:5])}', 'error')
                
    except Exception as e:
        error_msg = f'Error processing file: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.students'))

@management_bp.route('/students/<int:student_id>/toggle-status', methods=['POST'])
@login_required('management')
def toggle_student_status(student_id):
    """On deactivate: permanently delete the student; on inactive -> active: reactivate."""
    try:
        student = Student.query.get_or_404(student_id)

        if student.is_active:
            success, message = ManagementService.delete_student_permanently(student_id)
            if not success:
                if is_ajax_request():
                    return jsonify({'success': False, 'message': message})
                flash(message, 'error')
                return redirect(url_for('management.students'))

            if is_ajax_request():
                return jsonify({'success': True, 'message': message})
            flash(message, 'success')
            return redirect(url_for('management.students'))
        else:
            student.is_active = True
            db.session.commit()
            message = f'Student {student.name} has been activated'
            if is_ajax_request():
                return jsonify({'success': True, 'message': message})
            flash(message, 'success')
            return redirect(url_for('management.students'))

    except Exception as e:
        error_msg = f'Error updating student status: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.students'))

@management_bp.route('/students/bulk-toggle-status', methods=['POST'])
@login_required('management')
def bulk_toggle_student_status():
    """Bulk toggle status for multiple students"""
    try:
        student_ids = request.form.getlist('student_ids[]')
        action = request.form.get('action')  # 'activate' or 'deactivate'
        
        if not student_ids:
            if is_ajax_request():
                return jsonify({'success': False, 'message': 'No students selected'})
            flash('No students selected', 'error')
            return redirect(url_for('management.students'))
        
        success_count = 0
        error_messages = []
        
        for student_id in student_ids:
            try:
                student = Student.query.get_or_404(int(student_id))
                
                if action == 'deactivate' and student.is_active:
                    success, message = ManagementService.delete_student_permanently(int(student_id))
                    if success:
                        success_count += 1
                    else:
                        error_messages.append(f"{student.name}: {message}")
                elif action == 'activate' and not student.is_active:
                    student.is_active = True
                    db.session.commit()
                    success_count += 1
                # Skip if action doesn't match current status
                
            except Exception as e:
                error_messages.append(f"Student ID {student_id}: {str(e)}")
        
        message = f'Successfully processed {success_count} students'
        if error_messages:
            message += f'. Errors: {"; ".join(error_messages[:3])}'
            if len(error_messages) > 3:
                message += f' (+{len(error_messages)-3} more)'
        
        if is_ajax_request():
            return jsonify({'success': success_count > 0, 'message': message})
        
        if success_count > 0:
            flash(message, 'success')
        else:
            flash('No students were processed', 'warning')
            
    except Exception as e:
        error_msg = f'Error processing bulk operation: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.students'))

@management_bp.route('/courses/add', methods=['POST'])
@login_required('management')
def add_course():
    """Add new course"""
    try:
        course_data = {
            'name': request.form.get('name', '').strip(),
            'code': request.form.get('code', '').strip(),
            'description': request.form.get('description', '').strip() or None,
            'duration_years': int(request.form.get('duration_years', 3)),
            'total_semesters': int(request.form.get('total_semesters', 6))
        }
        
        success, message = ManagementService.create_course(course_data)
        
        # Handle AJAX requests
        if is_ajax_request():
            return jsonify({'success': success, 'message': message})
        
        if success:
            flash(message, 'success')
        else:
            flash(message, 'error')
            
    except Exception as e:
        error_msg = f'Error adding course: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.courses'))

@management_bp.route('/courses/<int:course_id>/toggle-status', methods=['POST'])
@login_required('management')
def toggle_course_status(course_id):
    """On deactivate: permanently delete the course; on inactive -> active: reactivate."""
    try:
        course = Course.query.get_or_404(course_id)

        if course.is_active:
            success, message = ManagementService.delete_course_permanently(course_id)
            if not success:
                if is_ajax_request():
                    return jsonify({'success': False, 'message': message})
                flash(message, 'error')
                return redirect(url_for('management.courses'))

            if is_ajax_request():
                return jsonify({'success': True, 'message': message})
            flash(message, 'success')
            return redirect(url_for('management.courses'))
        else:
            course.is_active = True
            db.session.commit()
            message = f'Course {course.name} has been activated'
            if is_ajax_request():
                return jsonify({'success': True, 'message': message})
            flash(message, 'success')
            return redirect(url_for('management.courses'))

    except Exception as e:
        error_msg = f'Error updating course status: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.courses'))

@management_bp.route('/subjects/add', methods=['POST'])
@login_required('management')
def add_subject():
    """Add new subject"""
    try:
        subject_data = {
            'name': request.form.get('name', '').strip(),
            'code': request.form.get('code', '').strip(),
            'course_id': int(request.form.get('course_id')),
            'year': int(request.form.get('year')),
            'semester': int(request.form.get('semester')),
            'description': request.form.get('description', '').strip() or None
        }
        
        success, message = ManagementService.create_subject(subject_data)
        
        # Handle AJAX requests
        if is_ajax_request():
            return jsonify({'success': success, 'message': message})
        
        if success:
            flash(message, 'success')
        else:
            flash(message, 'error')
            
    except Exception as e:
        error_msg = f'Error adding subject: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.subjects'))

@management_bp.route('/subjects/<int:subject_id>/toggle-status', methods=['POST'])
@login_required('management')
def toggle_subject_status(subject_id):
    """On deactivate: permanently delete the subject; on inactive -> active: reactivate."""
    try:
        subject = Subject.query.get_or_404(subject_id)

        if subject.is_active:
            success, message = ManagementService.delete_subject_permanently(subject_id)
            if not success:
                if is_ajax_request():
                    return jsonify({'success': False, 'message': message})
                flash(message, 'error')
                return redirect(url_for('management.subjects'))

            if is_ajax_request():
                return jsonify({'success': True, 'message': message})
            flash(message, 'success')
            return redirect(url_for('management.subjects'))
        else:
            subject.is_active = True
            db.session.commit()
            message = f'Subject {subject.name} has been activated'
            if is_ajax_request():
                return jsonify({'success': True, 'message': message})
            flash(message, 'success')
            return redirect(url_for('management.subjects'))

    except Exception as e:
        error_msg = f'Error updating subject status: {str(e)}'
        if is_ajax_request():
            return jsonify({'success': False, 'message': error_msg})
        flash(error_msg, 'error')
    
    return redirect(url_for('management.subjects'))

# ============================================================================
# REPORTING ROUTES
# ============================================================================

@management_bp.route('/reports')
@login_required('management')
def reports_dashboard():
    """Reports dashboard"""
    try:
        subjects = ReportingService.get_subjects_for_reporting()
        courses = ReportingService.get_courses_for_reporting()
        students = ReportingService.get_students_for_reporting()
        
        return render_template('management/reports.html', 
                             subjects=subjects, 
                             courses=courses,
                             students=students)
    except Exception as e:
        flash(f'Error loading reports dashboard: {str(e)}', 'error')
        return render_template('management/reports.html', subjects=[], courses=[], students=[])
        return render_template('management/reports.html', subjects=[], courses=[], students=[])

@management_bp.route('/reports/students/filter')
@login_required('management')
def filter_students_for_reports():
    """Filter students by course for reports"""
    try:
        course_id = request.args.get('course_id', type=int)
        students = ReportingService.get_students_for_reporting(course_id)
        return jsonify({'success': True, 'students': students})
    except Exception as e:
        return jsonify({'success': False, 'message': f'Error filtering students: {str(e)}'})

@management_bp.route('/reports/student/<int:student_id>')
@login_required('management')
def student_report(student_id):
    """Individual student report"""
    try:
        report = ReportingService.get_student_detailed_report(student_id)
        if not report:
            flash('Student not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        
        return render_template('management/student_report.html', report=report)
    except Exception as e:
        flash(f'Error generating student report: {str(e)}', 'error')
        return redirect(url_for('management.reports_dashboard'))

@management_bp.route('/reports/class/marks/<int:subject_id>')
@login_required('management')
def class_marks_report(subject_id):
    """Class marks report"""
    try:
        assessment_type = request.args.get('assessment_type')
        report = ReportingService.get_class_marks_report(subject_id, assessment_type)
        
        if not report:
            flash('Subject not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        
        return render_template('management/class_marks_report.html', report=report)
    except Exception as e:
        flash(f'Error generating class marks report: {str(e)}', 'error')
        return redirect(url_for('management.reports_dashboard'))

@management_bp.route('/reports/class/attendance/<int:subject_id>')
@login_required('management')
def class_attendance_report(subject_id):
    """Class attendance report"""
    try:
        month = request.args.get('month', type=int)
        year = request.args.get('year', type=int)
        
        report = ReportingService.get_class_attendance_report(subject_id, month, year)
        
        if not report:
            flash('Subject not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        
        return render_template('management/class_attendance_report.html', report=report)
    except Exception as e:
        flash(f'Error generating class attendance report: {str(e)}', 'error')
        return redirect(url_for('management.reports_dashboard'))

@management_bp.route('/reports/course/<int:course_id>')
@login_required('management')
def course_overview_report(course_id):
    """Course overview report"""
    try:
        report = ReportingService.get_course_overview_report(course_id)
        
        if not report:
            flash('Course not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        
        return render_template('management/course_overview_report.html', report=report)
    except Exception as e:
        flash(f'Error generating course overview report: {str(e)}', 'error')
        return redirect(url_for('management.reports_dashboard'))

# ============================================================================
# EXCEL EXPORT ROUTES
# ============================================================================

@management_bp.route('/reports/export/student/<int:student_id>')
@login_required('management')
def export_student_report(student_id):
    """Export student report to Excel"""
    try:
        from flask import make_response
        
        report = ReportingService.get_student_detailed_report(student_id)
        if not report:
            flash('Student not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        
        workbook = ExcelExportService.export_student_report(report)
        if not workbook:
            flash('Error generating Excel file', 'error')
            return redirect(url_for('management.student_report', student_id=student_id))
        
        excel_data = ExcelExportService.workbook_to_bytes(workbook)
        
        response = make_response(excel_data)
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f'attachment; filename=student_report_{report["student"]["roll_number"]}.xlsx'
        
        return response
        
    except Exception as e:
        flash(f'Error exporting student report: {str(e)}', 'error')
        return redirect(url_for('management.student_report', student_id=student_id))

# ---------------- PDF Export ----------------
@management_bp.route('/reports/export/student/<int:student_id>/pdf')
@login_required('management')
def export_student_report_pdf(student_id):
    """Export student report to PDF"""
    try:
        from flask import make_response
        report = ReportingService.get_student_detailed_report(student_id)
        if not report:
            flash('Student not found', 'error')
            return redirect(url_for('management.reports_dashboard'))

        pdf_bytes = ReportingService.generate_student_report_pdf(report)
        response = make_response(pdf_bytes)
        response.headers['Content-Type'] = 'application/pdf'
        response.headers['Content-Disposition'] = f'attachment; filename=student_report_{report["student"]["roll_number"]}.pdf'
        return response
    except Exception as e:
        flash(f'Error exporting PDF: {str(e)}', 'error')
        return redirect(url_for('management.student_report', student_id=student_id))

@management_bp.route('/reports/export/class/marks/<int:subject_id>/pdf')
@login_required('management')
def export_class_marks_report_pdf(subject_id):
    """Export class marks report to PDF"""
    try:
        from flask import make_response
        assessment_type = request.args.get('assessment_type')
        report = ReportingService.get_class_marks_report(subject_id, assessment_type)
        if not report:
            flash('Subject not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        pdf_bytes = ReportingService.generate_class_marks_report_pdf(report)
        filename = f"class_marks_{report['subject']['code']}"
        if assessment_type:
            filename += f"_{assessment_type}"
        filename += ".pdf"
        response = make_response(pdf_bytes)
        response.headers['Content-Type'] = 'application/pdf'
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        return response
    except Exception as e:
        flash(f'Error exporting class marks PDF: {str(e)}', 'error')
        return redirect(url_for('management.class_marks_report', subject_id=subject_id))

@management_bp.route('/reports/export/class/attendance/<int:subject_id>/pdf')
@login_required('management')
def export_class_attendance_report_pdf(subject_id):
    """Export class attendance report to PDF"""
    try:
        from flask import make_response
        month = request.args.get('month', type=int)
        year = request.args.get('year', type=int)
        report = ReportingService.get_class_attendance_report(subject_id, month, year)
        if not report:
            flash('Subject not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        pdf_bytes = ReportingService.generate_class_attendance_report_pdf(report)
        filename = f"class_attendance_{report['subject']['code']}_{report['month']}_{report['year']}.pdf"
        response = make_response(pdf_bytes)
        response.headers['Content-Type'] = 'application/pdf'
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        return response
    except Exception as e:
        flash(f'Error exporting class attendance PDF: {str(e)}', 'error')
        return redirect(url_for('management.class_attendance_report', subject_id=subject_id))

@management_bp.route('/reports/export/course/<int:course_id>/pdf')
@login_required('management')
def export_course_overview_report_pdf(course_id):
    """Export course overview report to PDF"""
    try:
        from flask import make_response
        report = ReportingService.get_course_overview_report(course_id)
        if not report:
            flash('Course not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        pdf_bytes = ReportingService.generate_course_overview_report_pdf(report)
        filename = f"course_overview_{report['course']['code']}.pdf"
        response = make_response(pdf_bytes)
        response.headers['Content-Type'] = 'application/pdf'
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        return response
    except Exception as e:
        flash(f'Error exporting course overview PDF: {str(e)}', 'error')
        return redirect(url_for('management.course_overview_report', course_id=course_id))

@management_bp.route('/reports/export/class/marks/<int:subject_id>')
@login_required('management')
def export_class_marks_report(subject_id):
    """Export class marks report to Excel"""
    try:
        from flask import make_response
        
        assessment_type = request.args.get('assessment_type')
        report = ReportingService.get_class_marks_report(subject_id, assessment_type)
        
        if not report:
            flash('Subject not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        
        workbook = ExcelExportService.export_class_marks_report(report)
        if not workbook:
            flash('Error generating Excel file', 'error')
            return redirect(url_for('management.class_marks_report', subject_id=subject_id))
        
        excel_data = ExcelExportService.workbook_to_bytes(workbook)
        
        filename = f"class_marks_{report['subject']['code']}"
        if assessment_type:
            filename += f"_{assessment_type}"
        filename += ".xlsx"
        
        response = make_response(excel_data)
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        
        return response
        
    except Exception as e:
        flash(f'Error exporting class marks report: {str(e)}', 'error')
        return redirect(url_for('management.class_marks_report', subject_id=subject_id))

@management_bp.route('/reports/export/class/attendance/<int:subject_id>')
@login_required('management')
def export_class_attendance_report(subject_id):
    """Export class attendance report to Excel"""
    try:
        from flask import make_response
        
        month = request.args.get('month', type=int)
        year = request.args.get('year', type=int)
        
        report = ReportingService.get_class_attendance_report(subject_id, month, year)
        
        if not report:
            flash('Subject not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        
        workbook = ExcelExportService.export_class_attendance_report(report)
        if not workbook:
            flash('Error generating Excel file', 'error')
            return redirect(url_for('management.class_attendance_report', subject_id=subject_id))
        
        excel_data = ExcelExportService.workbook_to_bytes(workbook)
        
        filename = f"class_attendance_{report['subject']['code']}_{report['month']}_{report['year']}.xlsx"
        
        response = make_response(excel_data)
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        
        return response
        
    except Exception as e:
        flash(f'Error exporting class attendance report: {str(e)}', 'error')
        return redirect(url_for('management.class_attendance_report', subject_id=subject_id))

@management_bp.route('/reports/export/course/<int:course_id>')
@login_required('management')
def export_course_overview_report(course_id):
    """Export course overview report to Excel"""
    try:
        from flask import make_response
        
        report = ReportingService.get_course_overview_report(course_id)
        
        if not report:
            flash('Course not found', 'error')
            return redirect(url_for('management.reports_dashboard'))
        
        workbook = ExcelExportService.export_course_overview_report(report)
        if not workbook:
            flash('Error generating Excel file', 'error')
            return redirect(url_for('management.course_overview_report', course_id=course_id))
        
        excel_data = ExcelExportService.workbook_to_bytes(workbook)
        
        filename = f"course_overview_{report['course']['code']}.xlsx"
        
        response = make_response(excel_data)
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        
        return response
        
    except Exception as e:
        flash(f'Error exporting course overview report: {str(e)}', 'error')
        return redirect(url_for('management.course_overview_report', course_id=course_id))