"""
Reporting service for Moulya College Management System
Handles marks and attendance reports for management
"""

from models.student import Student
from models.academic import Course, Subject
from models.marks import StudentMarks
from models.attendance import AttendanceRecord, MonthlyAttendanceSummary
from models.user import Lecturer
from database import db
from datetime import datetime, date
from sqlalchemy import func, and_, or_
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from xml.sax.saxutils import escape as xml_escape
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ReportingService:
    """Service for generating reports"""
    
    @staticmethod
    def _get_paragraph_style():
        """Return a compact cell Paragraph style to enable auto word-wrap in table cells."""
        styles = getSampleStyleSheet()
        from reportlab.lib.styles import ParagraphStyle
        return ParagraphStyle(
            'Cell',
            parent=styles['Normal'],
            fontSize=9,
            leading=11,
            spaceAfter=0,
            spaceBefore=0,
        )

    @staticmethod
    def _get_header_paragraph_style():
        styles = getSampleStyleSheet()
        from reportlab.lib.styles import ParagraphStyle
        return ParagraphStyle(
            'HeaderCell',
            parent=styles['Normal'],
            fontSize=9,
            leading=11,
            textColor=colors.white,
            spaceAfter=0,
            spaceBefore=0,
        )

    @staticmethod
    def _to_paragraph(value):
        """Convert any value to a Paragraph so ReportLab wraps text within cell width."""
        try:
            if value is None:
                return Paragraph('', ReportingService._get_paragraph_style())
            # Keep numbers as plain strings too; Paragraph will handle fine
            text = xml_escape(str(value))
            return Paragraph(text, ReportingService._get_paragraph_style())
        except Exception:
            return Paragraph(xml_escape(str(value)), ReportingService._get_paragraph_style())

    @staticmethod
    def _wrap_table_data(rows, skip_header=True, header_text_white=False):
        """Map table cells to Paragraphs for word-wrap.
        If skip_header=True, the first row is left as-is so TableStyle header
        text color/background rules still apply.
        """
        if not rows:
            return rows
        start_idx = 1 if skip_header and len(rows) > 0 else 0
        wrapped_rows = []
        # Keep header as-is if requested
        if start_idx == 1:
            if header_text_white:
                wrapped_rows.append([Paragraph(xml_escape(str(c)), ReportingService._get_header_paragraph_style()) for c in rows[0]])
            else:
                wrapped_rows.append(rows[0])
        for row in rows[start_idx:]:
            wrapped_rows.append([ReportingService._to_paragraph(cell) for cell in row])
        return wrapped_rows
    
    @staticmethod
    def _full_width_colwidths(total_width, num_columns):
        """Return equal column widths that use the full available width.
        This keeps every table consistent and maximized to the page width.
        """
        if num_columns <= 0:
            return []
        col = total_width / float(num_columns)
        return [col] * num_columns

    @staticmethod
    def _parse_course_and_section(course_name: str):
        """Split a user-entered course string into course and section.
        Examples:
        - "III bca b" -> ("bca", "b")
        - "I bcom" -> ("bcom", None)
        - "bsc" -> ("bsc", None)
        We consider the last token as section only when there are at least 3 tokens.
        The course becomes the last token when there are 1-2 tokens, otherwise the second last.
        """
        if not course_name:
            return None, None
        parts = [p for p in course_name.strip().split(' ') if p]
        if not parts:
            return None, None
        if len(parts) >= 3:
            return parts[-2], parts[-1]
        # len == 1 or 2 -> course is last token, no section
        return parts[-1], None
    
    @staticmethod
    def get_student_detailed_report(student_id):
        """Get detailed report for a specific student"""
        try:
            student = Student.query.get(student_id)
            if not student:
                return None
            
            # Get all subjects the student is enrolled in
            from models.student import StudentEnrollment
            subjects = Subject.query.join(StudentEnrollment).filter(
                StudentEnrollment.student_id == student_id,
                StudentEnrollment.is_active == True
            ).order_by(Subject.name.asc(), Subject.code.asc()).all()
            
            course_display, section = ReportingService._parse_course_and_section(student.course.name if student.course else None)

            report = {
                'student': {
                    'id': student.id,
                    'name': student.name,
                    'roll_number': student.roll_number,
                    'course': student.course.name if student.course else None,
                    'course_display': course_display,
                    'section': section,
                    'academic_year': student.academic_year,
                    'current_semester': student.current_semester
                },
                'subjects': []
            }
            
            for subject in subjects:
                # Get marks for this subject
                from models.marks import StudentMarks
                marks = StudentMarks.query.filter_by(
                    student_id=student_id,
                    subject_id=subject.id
                ).all()
                
                # Get attendance for this subject
                attendance_records = AttendanceRecord.query.filter_by(
                    student_id=student_id,
                    subject_id=subject.id
                ).all()
                
                total_classes = len(attendance_records)
                present_classes = len([r for r in attendance_records if r.status == 'present'])
                attendance_percentage = round((present_classes / total_classes) * 100, 2) if total_classes > 0 else 0
                
                # Calculate overall marks percentage for this subject
                if marks:
                    total_obtained = sum(mark.marks_obtained for mark in marks)
                    total_max = sum(mark.max_marks for mark in marks)
                    overall_percentage = round((total_obtained / total_max) * 100, 2) if total_max > 0 else 0
                else:
                    overall_percentage = 0
                
                # Format marks data with proper assessment type names
                marks_data = []
                for mark in marks:
                    mark_dict = mark.to_dict()
                    # Ensure assessment_type is properly formatted
                    assessment_type = mark.assessment_type
                    if assessment_type == 'internal1':
                        mark_dict['assessment_type'] = 'Internal 1'
                    elif assessment_type == 'internal2':
                        mark_dict['assessment_type'] = 'Internal 2'
                    elif assessment_type == 'assignment':
                        mark_dict['assessment_type'] = 'Assignment'
                    elif assessment_type == 'project':
                        mark_dict['assessment_type'] = 'Project'
                    else:
                        mark_dict['assessment_type'] = assessment_type.title()
                    
                    # Add percentage calculation
                    if mark.max_marks > 0:
                        mark_dict['percentage'] = round((mark.marks_obtained / mark.max_marks) * 100, 2)
                    else:
                        mark_dict['percentage'] = 0
                    
                    marks_data.append(mark_dict)
                
                subject_data = {
                    'subject_id': subject.id,
                    'subject_name': subject.name,
                    'subject_code': subject.code,
                    'year': subject.year,
                    'semester': subject.semester,
                    'marks': marks_data,
                    'overall_marks_percentage': overall_percentage,
                    'attendance': {
                        'total_classes': total_classes,
                        'present_classes': present_classes,
                        'absent_classes': total_classes - present_classes,
                        'attendance_percentage': attendance_percentage
                    }
                }
                
                report['subjects'].append(subject_data)
            
            return report
            
        except Exception as e:
            print(f"Error generating student report: {e}")
            return None
    
    @staticmethod
    def get_class_marks_report(subject_id, assessment_type=None):
        """Get marks report for entire class in a subject"""
        try:
            subject = Subject.query.get(subject_id)
            if not subject:
                return None
            
            from models.marks import StudentMarks
            query = StudentMarks.query.filter_by(subject_id=subject_id)
            if assessment_type:
                query = query.filter_by(assessment_type=assessment_type)
            
            marks = query.all()
            
            # Group marks by student
            student_marks = {}
            for mark in marks:
                student_id = mark.student_id
                if student_id not in student_marks:
                    student_marks[student_id] = {
                        'student': mark.student,
                        'student_name': mark.student.name,
                        'roll_number': mark.student.roll_number,
                        'marks': []
                    }
                
                mark_dict = mark.to_dict()
                # Format assessment type for display
                assessment_type_display = mark.assessment_type
                if assessment_type_display == 'internal1':
                    mark_dict['assessment_type'] = 'Internal 1'
                elif assessment_type_display == 'internal2':
                    mark_dict['assessment_type'] = 'Internal 2'
                elif assessment_type_display == 'assignment':
                    mark_dict['assessment_type'] = 'Assignment'
                elif assessment_type_display == 'project':
                    mark_dict['assessment_type'] = 'Project'
                else:
                    mark_dict['assessment_type'] = assessment_type_display.title()
                
                student_marks[student_id]['marks'].append(mark_dict)
            
            # Calculate statistics
            all_percentages = []
            for mark in marks:
                if mark.max_marks > 0:
                    percentage = round((mark.marks_obtained / mark.max_marks) * 100, 2)
                    all_percentages.append(percentage)
            
            statistics = {
                'total_students': len(student_marks),
                'total_assessments': len(marks),
                'class_average': round(sum(all_percentages) / len(all_percentages), 2) if all_percentages else 0,
                'highest_score': max(all_percentages) if all_percentages else 0,
                'lowest_score': min(all_percentages) if all_percentages else 0,
                'passing_students': len([p for p in all_percentages if p >= 35]),
                'failing_students': len([p for p in all_percentages if p < 35])
            }
            
            # Format assessment type for display
            assessment_type_display = assessment_type
            if assessment_type == 'internal1':
                assessment_type_display = 'Internal 1'
            elif assessment_type == 'internal2':
                assessment_type_display = 'Internal 2'
            elif assessment_type == 'assignment':
                assessment_type_display = 'Assignment'
            elif assessment_type == 'project':
                assessment_type_display = 'Project'
            elif assessment_type:
                assessment_type_display = assessment_type.title()
            
            course_display, section = ReportingService._parse_course_and_section(subject.course.name if subject.course else None)

            lecturers_list = []
            try:
                lecturers_list = [lec.name for lec in subject.get_assigned_lecturers() if lec]
            except Exception:
                lecturers_list = []

            # Sort student buckets by roll number then name (ascending)
            sorted_student_values = sorted(
                student_marks.values(),
                key=lambda x: ((x.get('roll_number') or '').upper(), (x.get('student_name') or '').upper())
            )

            report = {
                'subject': {
                    'id': subject.id,
                    'name': subject.name,
                    'code': subject.code,
                    'course': subject.course.name if subject.course else None,
                    'course_display': course_display,
                    'section': section,
                    'year': subject.year,
                    'semester': subject.semester,
                    'lecturers': lecturers_list
                },
                'assessment_type': assessment_type_display,
                'statistics': statistics,
                'student_marks': sorted_student_values
            }
            
            return report
            
        except Exception as e:
            print(f"Error generating class marks report: {e}")
            return None
    
    @staticmethod
    def get_class_attendance_report(subject_id, month=None, year=None):
        """Get attendance report for entire class in a subject"""
        try:
            subject = Subject.query.get(subject_id)
            if not subject:
                return None
            
            # Default to current month/year if not specified
            if not month:
                month = datetime.now().month
            if not year:
                year = datetime.now().year
            
            # Get all students enrolled in this subject
            from models.student import StudentEnrollment
            students = Student.query.join(StudentEnrollment).filter(
                StudentEnrollment.subject_id == subject_id,
                StudentEnrollment.is_active == True
            ).order_by(Student.roll_number.asc(), Student.name.asc()).all()
            
            student_attendance = []
            total_classes_conducted = 0
            
            # First, get the total number of classes conducted for this subject in the month
            all_attendance_records = AttendanceRecord.query.filter(
                AttendanceRecord.subject_id == subject_id,
                db.extract('month', AttendanceRecord.date) == month,
                db.extract('year', AttendanceRecord.date) == year
            ).all()
            
            # Get unique dates to count total classes
            unique_dates = set(record.date for record in all_attendance_records)
            total_classes_conducted = len(unique_dates)
            
            from datetime import date as date_class
            import calendar
            prev_year = year
            prev_month = month - 1
            if prev_month == 0:
                prev_month = 12
                prev_year -= 1

            for student in students:
                # cumulative present till selected month
                month_end = date_class(year, month, calendar.monthrange(year, month)[1])
                prev_end = date_class(prev_year, prev_month, calendar.monthrange(prev_year, prev_month)[1])

                cum_present_till_month = AttendanceRecord.query.filter(
                    AttendanceRecord.student_id == student.id,
                    AttendanceRecord.subject_id == subject_id,
                    AttendanceRecord.date <= month_end,
                    AttendanceRecord.status == 'present'
                ).count()

                cum_present_till_prev = AttendanceRecord.query.filter(
                    AttendanceRecord.student_id == student.id,
                    AttendanceRecord.subject_id == subject_id,
                    AttendanceRecord.date <= prev_end,
                    AttendanceRecord.status == 'present'
                ).count()

                present_classes = max(cum_present_till_month - cum_present_till_prev, 0)
                total_classes = total_classes_conducted
                absent_classes = max(total_classes_conducted - present_classes, 0)
                
                # Calculate percentage based on delta present vs conducted this month
                attendance_percentage = round((present_classes / total_classes) * 100, 2) if total_classes > 0 else 0
                
                student_data = {
                    'student_id': student.id,
                    'student_name': student.name,
                    'roll_number': student.roll_number,
                    'total_classes': total_classes_conducted,  # Use consistent total
                    'present_classes': present_classes,
                    'absent_classes': total_classes_conducted - present_classes,
                    'attendance_percentage': attendance_percentage,
                    'status': 'Good' if attendance_percentage >= 75 else 'Poor' if attendance_percentage < 50 else 'Average'
                }
                
                student_attendance.append(student_data)
            
            # Calculate class statistics
            all_percentages = [s['attendance_percentage'] for s in student_attendance]
            valid_percentages = [p for p in all_percentages if p > 0]
            
            statistics = {
                'total_students': len(students),
                'total_classes_conducted': total_classes_conducted,
                'class_average_attendance': round(sum(valid_percentages) / len(valid_percentages), 2) if valid_percentages else 0,
                'students_with_good_attendance': len([s for s in student_attendance if s['attendance_percentage'] >= 75]),
                'students_with_poor_attendance': len([s for s in student_attendance if s['attendance_percentage'] < 50])
            }
            
            # Get month name for display
            month_names = ['', 'January', 'February', 'March', 'April', 'May', 'June',
                          'July', 'August', 'September', 'October', 'November', 'December']
            month_name = month_names[month] if 1 <= month <= 12 else str(month)
            
            course_display, section = ReportingService._parse_course_and_section(subject.course.name if subject.course else None)

            lecturers_list = []
            try:
                lecturers_list = [lec.name for lec in subject.get_assigned_lecturers() if lec]
            except Exception:
                lecturers_list = []

            # Ensure deterministic ascending order by roll then name
            student_attendance_sorted = sorted(
                student_attendance,
                key=lambda s: ((s.get('roll_number') or '').upper(), (s.get('student_name') or '').upper())
            )

            report = {
                'subject': {
                    'id': subject.id,
                    'name': subject.name,
                    'code': subject.code,
                    'course': subject.course.name if subject.course else None,
                    'course_display': course_display,
                    'section': section,
                    'year': subject.year,
                    'semester': subject.semester,
                    'lecturers': lecturers_list
                },
                'month': month_name,
                'year': year,
                'statistics': statistics,
                'student_attendance': student_attendance_sorted
            }
            
            return report
            
        except Exception as e:
            print(f"Error generating class attendance report: {e}")
            return None
    
    @staticmethod
    def get_course_overview_report(course_id):
        """Get overview report for entire course"""
        try:
            course = Course.query.get(course_id)
            if not course:
                return None
            
            # Get all subjects in this course
            subjects = (Subject.query
                .filter_by(course_id=course_id, is_active=True)
                .order_by(Subject.code.asc(), Subject.name.asc())
                .all())
            
            # Get all students in this course
            students = Student.query.filter_by(course_id=course_id, is_active=True).order_by(Student.roll_number.asc(), Student.name.asc()).all()
            
            subject_reports = []
            for subject in subjects:
                # Get marks statistics
                marks = StudentMarks.query.filter_by(subject_id=subject.id).all()
                marks_percentages = [mark.percentage for mark in marks]
                
                # Get attendance statistics
                attendance_records = AttendanceRecord.query.filter_by(subject_id=subject.id).all()
                total_records = len(attendance_records)
                present_records = len([r for r in attendance_records if r.status == 'present'])
                
                subject_data = {
                    'subject_id': subject.id,
                    'subject_name': subject.name,
                    'subject_code': subject.code,
                    'year': subject.year,
                    'semester': subject.semester,
                    'enrolled_students': subject.get_enrolled_students_count(),
                    'marks_statistics': {
                        'total_assessments': len(marks),
                        'average_marks': round(sum(marks_percentages) / len(marks_percentages), 2) if marks_percentages else 0,
                        'passing_rate': round((len([p for p in marks_percentages if p >= 35]) / len(marks_percentages)) * 100, 2) if marks_percentages else 0
                    },
                    'attendance_statistics': {
                        'total_classes': total_records,
                        'average_attendance': round((present_records / total_records) * 100, 2) if total_records > 0 else 0
                    }
                }
                
                subject_reports.append(subject_data)
            
            # Sort subjects in the overview report (by code, then name)
            subject_reports_sorted = sorted(
                subject_reports,
                key=lambda s: ((s.get('subject_code') or '').upper(), (s.get('subject_name') or '').upper())
            )

            report = {
                'course': {
                    'id': course.id,
                    'name': course.name,
                    'code': course.code,
                    'duration_years': course.duration_years,
                    'total_semesters': course.total_semesters
                },
                'total_students': len(students),
                'total_subjects': len(subjects),
                'subjects': subject_reports_sorted
            }
            
            return report
            
        except Exception as e:
            print(f"Error generating course overview report: {e}")
            return None
    
    @staticmethod
    def get_subjects_for_reporting():
        """Get all subjects available for reporting"""
        try:
            subjects = Subject.query.filter_by(is_active=True).order_by(Subject.name.asc(), Subject.code.asc()).all()
            return [{
                'id': subject.id,
                'name': subject.name,
                'code': subject.code,
                'course_name': subject.course.name if subject.course else None,
                'year': subject.year,
                'semester': subject.semester,
                'enrolled_students_count': subject.get_enrolled_students_count()
            } for subject in subjects]
        except Exception as e:
            print(f"Error getting subjects for reporting: {e}")
            return []
    
    @staticmethod
    def get_courses_for_reporting():
        """Get all courses available for reporting"""
        try:
            courses = Course.query.filter_by(is_active=True).order_by(Course.name.asc(), Course.code.asc()).all()
            return [{
                'id': course.id,
                'name': course.name,
                'code': course.code,
                'total_students': course.get_active_students_count(),
                'total_subjects': course.subjects.filter_by(is_active=True).count()
            } for course in courses]
        except Exception as e:
            print(f"Error getting courses for reporting: {e}")
            return []
    
    @staticmethod
    def get_students_for_reporting(course_id=None):
        """Get all students available for reporting, optionally filtered by course"""
        try:
            query = Student.query.filter_by(is_active=True)
            if course_id:
                query = query.filter_by(course_id=course_id)
            
            students = query.order_by(Student.roll_number.asc(), Student.name.asc()).all()
            return [{
                'id': student.id,
                'name': student.name,
                'roll_number': student.roll_number,
                'course_name': student.course.name if student.course else None,
                'course_id': student.course_id,
                'academic_year': student.academic_year,
                'current_semester': student.current_semester
            } for student in students]
        except Exception as e:
            print(f"Error getting students for reporting: {e}")
            return []

    # ======================== PDF GENERATION ========================
    @staticmethod
    def generate_student_report_pdf(report):
        """Generate a PDF for the student detailed report and return bytes."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
        elements = []
        styles = getSampleStyleSheet()
        from reportlab.lib.styles import ParagraphStyle
        title_center = ParagraphStyle('TitleCenter', parent=styles['Title'], alignment=1)
        subtitle_center = ParagraphStyle('SubtitleCenter', parent=styles['Normal'], alignment=1)
        header_title = ParagraphStyle('HeaderTitle', parent=styles['Title'], alignment=0, fontSize=16, leading=19)
        header_sub = ParagraphStyle('HeaderSub', parent=styles['Normal'], alignment=0, fontSize=10, leading=12)

        # Header with logo and college name
        # College-style header (logo + text, underline)
        try:
            from flask import current_app
            logo_path = current_app.root_path + '/static/img/logo-removebg-preview.png'
            logo_img = Image(logo_path)
            logo_img._restrictSize(26*mm, 26*mm)
        except Exception:
            logo_img = ''
        header_text = [
            Paragraph('Dr. B. B. Hegde First Grade College, Kundapura', header_title),
            Paragraph('A Unit of Coondapur Education Society (R)', header_sub)
        ]
        header_table = Table([[logo_img, header_text]], colWidths=[26*mm, (A4[0] - (18*mm + 18*mm) - 26*mm)])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LINEBELOW', (0,0), (-1,0), 0.75, colors.lightgrey),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(header_table)

        # Main report title
        elements.append(Spacer(1, 6))
        elements.append(Paragraph('Student Performance Report', title_center))
        elements.append(Spacer(1, 8))

        # Student info table
        s = report['student']
        course_line = s.get('course_display') or s.get('course') or ''
        data = [
            ['Name', s.get('name', '')],
            ['Roll Number', s.get('roll_number', '')],
            ['Course', course_line],
        ]
        if s.get('section'):
            data.append(['Section', str(s.get('section')).upper()])
        data.extend([
            ['Academic Year', s.get('academic_year', '')],
            ['Current Semester', s.get('current_semester', '')]
        ])

        table = Table(data, colWidths=[45*mm, 115*mm])
        table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
            # no header row now
            ('ALIGN', (0,0), (0,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        elements.extend([Spacer(1, 6), table, Spacer(1, 12)])

        # Marks table
        elements.append(Paragraph('Marks Report', styles['Heading2']))
        marks_headers = ['Subject', 'Code', 'Assessment', 'Marks', 'Max', 'Percent', 'Grade', 'Status']
        marks_rows = [marks_headers]
        for subj in report.get('subjects', []):
            for m in subj.get('marks', []):
                marks_rows.append([
                    subj.get('subject_name',''), subj.get('subject_code',''), m.get('assessment_type',''),
                    m.get('marks_obtained',''), m.get('max_marks',''), m.get('percentage',''), m.get('grade',''), m.get('performance_status','')
                ])
        if len(marks_rows) == 1:
            marks_rows.append(['No data'] + ['']*7)
        # Wrap text in table data and make table full width
        marks_rows_wrapped = ReportingService._wrap_table_data(marks_rows, skip_header=True, header_text_white=True)
        page_width = A4[0] - (18*mm + 18*mm)
        marks_table = Table(marks_rows_wrapped, repeatRows=1, colWidths=ReportingService._full_width_colwidths(page_width, len(marks_headers)))
        marks_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
            ('BACKGROUND', (0,0), (-1,0), colors.black),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('ALIGN', (0,0), (-1,0), 'CENTER'),
            ('VALIGN', (0,0), (-1,0), 'MIDDLE'),
            ('LEFTPADDING', (0,0), (-1,0), 4),
            ('RIGHTPADDING', (0,0), (-1,0), 4),
            ('TOPPADDING', (0,0), (-1,0), 3),
            ('BOTTOMPADDING', (0,0), (-1,0), 3)
        ]))
        elements.extend([marks_table, Spacer(1, 12)])

        # Attendance table
        elements.append(Paragraph('Attendance Report', styles['Heading2']))
        att_headers = ['Subject', 'Code', 'Total', 'Present', 'Absent', 'Percent', 'Status']
        att_rows = [att_headers]
        for subj in report.get('subjects', []):
            a = subj.get('attendance', {})
            att_rows.append([
                subj.get('subject_name',''), subj.get('subject_code',''), a.get('total_classes',''),
                a.get('present_classes',''), a.get('absent_classes',''), a.get('attendance_percentage',''),
                'Good' if a.get('attendance_percentage',0) >= 75 else 'Average' if a.get('attendance_percentage',0) >= 50 else 'Poor'
            ])
        if len(att_rows) == 1:
            att_rows.append(['No data'] + ['']*6)
        # Wrap text in table data
        att_rows_wrapped = ReportingService._wrap_table_data(att_rows, skip_header=True, header_text_white=True)
        att_table = Table(att_rows_wrapped, repeatRows=1, colWidths=ReportingService._full_width_colwidths(page_width, len(att_headers)))
        att_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
            ('BACKGROUND', (0,0), (-1,0), colors.black),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold')
        ]))
        elements.append(att_table)

        doc.build(elements)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    # ======================== LECTURER PDFS ========================
    @staticmethod
    def generate_subject_marks_report_pdf(subject, marks_report):
        """Generate a PDF for a subject's marks report (lecturer view)."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
        elements = []
        styles = getSampleStyleSheet()
        from reportlab.lib.styles import ParagraphStyle
        header_title = ParagraphStyle('HeaderTitle', parent=styles['Title'], alignment=0, fontSize=16, leading=19)
        header_sub = ParagraphStyle('HeaderSub', parent=styles['Normal'], alignment=0, fontSize=10, leading=12)

        # Header (logo + college text)
        try:
            from flask import current_app
            logo_path = current_app.root_path + '/static/img/logo-removebg-preview.png'
            logo_img = Image(logo_path)
            logo_img._restrictSize(26*mm, 26*mm)
        except Exception:
            logo_img = ''
        header_text = [
            Paragraph('Dr. B. B. Hegde First Grade College, Kundapura', header_title),
            Paragraph('A Unit of Coondapur Education Society (R)', header_sub)
        ]
        header_table = Table([[logo_img, header_text]], colWidths=[26*mm, (A4[0] - (18*mm + 18*mm) - 26*mm)])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LINEBELOW', (0,0), (-1,0), 0.75, colors.lightgrey),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(header_table)

        # Get faculty name
        from models.assignments import SubjectAssignment
        assignment = SubjectAssignment.query.filter_by(subject_id=subject.id, is_active=True).first()
        faculty_name = assignment.lecturer.name if assignment and assignment.lecturer else 'N/A'
        
        # Parse course name to extract course code and section
        course_name = subject.course.name if subject.course else 'N/A'
        course_parts = course_name.split() if course_name != 'N/A' else []
        
        # Determine if section exists (if there are 3+ parts, last part is section)
        has_section = len(course_parts) >= 3
        course_code = course_parts[-2] if has_section else (course_parts[-1] if course_parts else 'N/A')
        section = course_parts[-1] if has_section else None
        
        # Subject summary box
        subj_rows = [
            ['Subject', subject.name],
            ['Code', subject.code],
            ['Course', course_code]
        ]
        if has_section:
            subj_rows.append(['Section', section])
        subj_rows.append(['Faculty', faculty_name])
        
        subj_table = Table(subj_rows, colWidths=[35*mm, (A4[0] - (18*mm + 18*mm) - 35*mm)])
        subj_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        elements.extend([Spacer(1, 10), Paragraph('Marks Report', styles['Heading2']), Spacer(1, 6), subj_table, Spacer(1, 10)])

        # Check if Internal assessments are available
        has_internal1 = any(record.get('internal1_marks') for record in marks_report or [])
        has_internal2 = any(record.get('internal2_marks') for record in marks_report or [])
        
        # Build table headers dynamically
        headers = ['Student', 'Roll Number']
        if has_internal1:
            headers.append('Internal 1')
        if has_internal2:
            headers.append('Internal 2')
        headers.extend(['Overall %', 'Status'])
        
        # Marks table
        rows = [headers]
        for record in marks_report or []:
            student = record['student']
            overall = record.get('overall_percentage') or 0
            status = 'Good' if overall >= 50 else 'Deficient'
            
            row = [student.name, student.roll_number]
            
            if has_internal1:
                internal1 = record.get('internal1_marks')
                if internal1:
                    row.append(f"{internal1['obtained']}/{internal1['max']}")
                else:
                    row.append('-')
            
            if has_internal2:
                internal2 = record.get('internal2_marks')
                if internal2:
                    row.append(f"{internal2['obtained']}/{internal2['max']}")
                else:
                    row.append('-')
            
            row.extend([f"{overall}%", status])
            rows.append(row)
            
        if len(rows) == 1:
            no_data_row = ['No data'] + [''] * (len(headers) - 1)
            rows.append(no_data_row)
            
        # Wrap text in table data
        rows_wrapped = ReportingService._wrap_table_data(rows, skip_header=True, header_text_white=True)

        page_width = A4[0] - (18*mm + 18*mm)
        table = Table(rows_wrapped, repeatRows=1, colWidths=ReportingService._full_width_colwidths(page_width, len(rows[0])))
        table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
            ('BACKGROUND', (0,0), (-1,0), colors.black),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold')
        ]))
        elements.append(table)

        doc.build(elements)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    @staticmethod
    def generate_subject_attendance_report_pdf(subject, attendance_report):
        """Generate a PDF for a subject's attendance report (lecturer view)."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
        elements = []
        styles = getSampleStyleSheet()
        from reportlab.lib.styles import ParagraphStyle
        header_title = ParagraphStyle('HeaderTitle', parent=styles['Title'], alignment=0, fontSize=16, leading=19)
        header_sub = ParagraphStyle('HeaderSub', parent=styles['Normal'], alignment=0, fontSize=10, leading=12)

        try:
            from flask import current_app
            logo_path = current_app.root_path + '/static/img/logo-removebg-preview.png'
            logo_img = Image(logo_path)
            logo_img._restrictSize(26*mm, 26*mm)
        except Exception:
            logo_img = ''
        header_text = [
            Paragraph('Dr. B. B. Hegde First Grade College, Kundapura', header_title),
            Paragraph('A Unit of Coondapur Education Society (R)', header_sub)
        ]
        header_table = Table([[logo_img, header_text]], colWidths=[26*mm, 148*mm])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LINEBELOW', (0,0), (-1,0), 0.75, colors.lightgrey),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(header_table)

        # Get faculty name
        from models.assignments import SubjectAssignment
        assignment = SubjectAssignment.query.filter_by(subject_id=subject.id, is_active=True).first()
        faculty_name = assignment.lecturer.name if assignment and assignment.lecturer else 'N/A'
        
        # Parse course name to extract course code and section
        course_name = subject.course.name if subject.course else 'N/A'
        course_parts = course_name.split() if course_name != 'N/A' else []
        
        # Determine if section exists (if there are 3+ parts, last part is section)
        has_section = len(course_parts) >= 3
        course_code = course_parts[-2] if has_section else (course_parts[-1] if course_parts else 'N/A')
        section = course_parts[-1] if has_section else None
        
        # Subject summary box
        subj_rows = [
            ['Subject', subject.name],
            ['Code', subject.code],
            ['Course', course_code]
        ]
        if has_section:
            subj_rows.append(['Section', section])
        subj_rows.append(['Faculty', faculty_name])
        
        subj_table = Table(subj_rows, colWidths=[35*mm, (A4[0] - (18*mm + 18*mm) - 35*mm)])
        subj_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        elements.extend([Spacer(1, 10), Paragraph('Attendance Report', styles['Heading2']), Spacer(1, 6), subj_table, Spacer(1, 10)])

        # Attendance table (Student | Present | Total | % | Status)
        rows = [['Student', 'Roll Number', 'Present', 'Total', '%', 'Status']]
        for record in attendance_report or []:
            student = record['student']
            percent = record.get('attendance_percentage') or 0
            status = 'Good' if percent >= 75 else 'Shortage'
            rows.append([
                student.name,
                student.roll_number,
                record.get('present_classes') or 0,
                record.get('total_classes') or 0,
                f"{percent}%",
                status
            ])
        if len(rows) == 1:
            rows.append(['No data', '', '', '', '', ''])
        # Wrap text in table data
        rows_wrapped = ReportingService._wrap_table_data(rows, skip_header=True, header_text_white=True)

        page_width = A4[0] - (18*mm + 18*mm)
        table = Table(rows_wrapped, repeatRows=1, colWidths=ReportingService._full_width_colwidths(page_width, len(rows[0])))
        table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
            ('BACKGROUND', (0,0), (-1,0), colors.black),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold')
        ]))
        elements.append(table)

        doc.build(elements)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    # ======================== EXCEL EXPORT FUNCTIONS ========================
    @staticmethod
    def generate_subject_marks_report_excel(subject, marks_report):
        """Generate an Excel file for a subject's marks report (lecturer view)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Marks Report"
        
        # Get faculty name
        from models.assignments import SubjectAssignment
        assignment = SubjectAssignment.query.filter_by(subject_id=subject.id, is_active=True).first()
        faculty_name = assignment.lecturer.name if assignment and assignment.lecturer else 'N/A'
        
        # Parse course name to extract course code and section
        course_name = subject.course.name if subject.course else 'N/A'
        course_parts = course_name.split() if course_name != 'N/A' else []
        
        # Determine if section exists (if there are 3+ parts, last part is section)
        has_section = len(course_parts) >= 3
        course_code = course_parts[-2] if has_section else (course_parts[-1] if course_parts else 'N/A')
        section = course_parts[-1] if has_section else None
        
        # Title
        ws['A1'] = 'Marks Report'
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:D1')
        
        # Subject info
        ws['A3'] = 'Subject'
        ws['B3'] = subject.name
        ws['A4'] = 'Code'
        ws['B4'] = subject.code
        ws['A5'] = 'Course'
        ws['B5'] = course_code
        if has_section:
            ws['A6'] = 'Section'
            ws['B6'] = section
            ws['A7'] = 'Faculty'
            ws['B7'] = faculty_name
        else:
            ws['A6'] = 'Faculty'
            ws['B6'] = faculty_name
        # Check if Internal assessments are available
        has_internal1 = any(record.get('internal1_marks') for record in marks_report or [])
        has_internal2 = any(record.get('internal2_marks') for record in marks_report or [])
        
        # Build table headers dynamically
        headers = ['Student', 'Roll Number']
        if has_internal1:
            headers.append('Internal 1')
        if has_internal2:
            headers.append('Internal 2')
        headers.extend(['Overall %', 'Status'])
        
        # Determine where to place spacer row and headers
        last_info_row = 7 if has_section else 6
        spacer_row = last_info_row + 1
        header_row = spacer_row + 1
        data_start_row = header_row + 1

        # Spacer row above the table
        ws.merge_cells(start_row=spacer_row, start_column=1, end_row=spacer_row, end_column=len(headers))
        spacer_cell = ws.cell(row=spacer_row, column=1, value='')
        spacer_cell.fill = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')

        # Table headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='000000', end_color='000000', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        # Data rows
        row = data_start_row
        for record in marks_report or []:
            student = record['student']
            overall = record.get('overall_percentage') or 0
            status = 'Good' if overall >= 50 else 'Deficient'
            
            col = 1
            ws.cell(row=row, column=col, value=student.name)
            col += 1
            ws.cell(row=row, column=col, value=student.roll_number)
            col += 1
            
            if has_internal1:
                internal1 = record.get('internal1_marks')
                if internal1:
                    ws.cell(row=row, column=col, value=f"{internal1['obtained']}/{internal1['max']}")
                else:
                    ws.cell(row=row, column=col, value='-')
                col += 1
            
            if has_internal2:
                internal2 = record.get('internal2_marks')
                if internal2:
                    ws.cell(row=row, column=col, value=f"{internal2['obtained']}/{internal2['max']}")
                else:
                    ws.cell(row=row, column=col, value='-')
                col += 1
            
            ws.cell(row=row, column=col, value=f"{overall}%")
            col += 1
            ws.cell(row=row, column=col, value=status)
            row += 1
        
        if not marks_report:
            ws.cell(row=data_start_row, column=1, value='No data')
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = None
            for cell in column:
                try:
                    if hasattr(cell, 'column_letter'):
                        column_letter = cell.column_letter
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            if column_letter:
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save to BytesIO
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_subject_attendance_report_excel(subject, attendance_report):
        """Generate an Excel file for a subject's attendance report (lecturer view)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Report"
        
        # Get faculty name
        from models.assignments import SubjectAssignment
        assignment = SubjectAssignment.query.filter_by(subject_id=subject.id, is_active=True).first()
        faculty_name = assignment.lecturer.name if assignment and assignment.lecturer else 'N/A'
        
        # Parse course name to extract course code and section
        course_name = subject.course.name if subject.course else 'N/A'
        course_parts = course_name.split() if course_name != 'N/A' else []
        
        # Determine if section exists (if there are 3+ parts, last part is section)
        has_section = len(course_parts) >= 3
        course_code = course_parts[-2] if has_section else (course_parts[-1] if course_parts else 'N/A')
        section = course_parts[-1] if has_section else None
        
        # Title
        ws['A1'] = 'Attendance Report'
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:F1')
        
        # Subject info
        ws['A3'] = 'Subject'
        ws['B3'] = subject.name
        ws['A4'] = 'Code'
        ws['B4'] = subject.code
        ws['A5'] = 'Course'
        ws['B5'] = course_code
        if has_section:
            ws['A6'] = 'Section'
            ws['B6'] = section
            ws['A7'] = 'Faculty'
            ws['B7'] = faculty_name
        else:
            ws['A6'] = 'Faculty'
            ws['B6'] = faculty_name
        # Determine where to place spacer row and headers
        last_info_row = 7 if has_section else 6
        spacer_row = last_info_row + 1
        header_row = spacer_row + 1
        data_start_row = header_row + 1

        # Spacer row above the table
        ws.merge_cells(start_row=spacer_row, start_column=1, end_row=spacer_row, end_column=6)
        spacer_cell = ws.cell(row=spacer_row, column=1, value='')
        spacer_cell.fill = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')

        # Table headers
        headers = ['Student', 'Roll Number', 'Present', 'Total', '%', 'Status']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='000000', end_color='000000', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        # Data rows
        row = data_start_row
        for record in attendance_report or []:
            student = record['student']
            percent = record.get('attendance_percentage') or 0
            status = 'Good' if percent >= 75 else 'Shortage'
            
            ws.cell(row=row, column=1, value=student.name)
            ws.cell(row=row, column=2, value=student.roll_number)
            ws.cell(row=row, column=3, value=record.get('present_classes') or 0)
            ws.cell(row=row, column=4, value=record.get('total_classes') or 0)
            ws.cell(row=row, column=5, value=f"{percent}%")
            ws.cell(row=row, column=6, value=status)
            row += 1
        
        if not attendance_report:
            ws.cell(row=data_start_row, column=1, value='No data')
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = None
            for cell in column:
                try:
                    if hasattr(cell, 'column_letter'):
                        column_letter = cell.column_letter
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            if column_letter:
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save to BytesIO
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    # ======================== ADDITIONAL PDF GENERATORS ========================
    @staticmethod
    def generate_class_marks_report_pdf(report):
        """Generate a PDF for the class marks report and return bytes."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
        elements = []
        styles = getSampleStyleSheet()
        # Header (same format as student PDF)
        try:
            from flask import current_app
            logo_path = current_app.root_path + '/static/img/logo-removebg-preview.png'
            logo_img = Image(logo_path)
            logo_img._restrictSize(26*mm, 26*mm)
        except Exception:
            logo_img = ''
        header_title = styles['Title']
        from reportlab.lib.styles import ParagraphStyle
        header_title = ParagraphStyle('HeaderTitle', parent=styles['Title'], alignment=0, fontSize=16, leading=19)
        header_sub = ParagraphStyle('HeaderSub', parent=styles['Normal'], alignment=0, fontSize=10, leading=12)
        header_text = [
            Paragraph('Dr. B. B. Hegde First Grade College, Kundapura', header_title),
            Paragraph('A Unit of Coondapur Education Society (R)', header_sub)
        ]
        header_table = Table([[logo_img, header_text]], colWidths=[26*mm, 148*mm])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LINEBELOW', (0,0), (-1,0), 0.75, colors.lightgrey),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(header_table)
        elements.append(Spacer(1, 6))
        elements.append(Paragraph('Class Marks Report', styles['Title']))
        s = report.get('subject', {})
        meta_rows = [
            ['Subject', f"{s.get('name','')} ({s.get('code','')})"],
            ['Course', s.get('course_display') or s.get('course') or '']
        ]
        if s.get('section'):
            meta_rows.append(['Section', str(s.get('section')).upper()])
        meta_rows.append(['Year/Sem', f"{s.get('year','')}/{s.get('semester','')}"])
        # Faculty line if available
        try:
            lecturers = s.get('lecturers') or []
            if lecturers:
                meta_rows.append(['Faculty', ', '.join(lecturers)])
        except Exception:
            pass
        if report.get('assessment_type'):
            meta_rows.append(['Assessment Type', report.get('assessment_type')])
        meta_table = Table(meta_rows, colWidths=[40*mm, 120*mm])
        meta_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
        ]))
        elements.extend([Spacer(1, 6), meta_table, Spacer(1, 8)])

        # Statistics
        stats = report.get('statistics', {})
        stats_rows = [
            ['Total Students', stats.get('total_students', 0)],
            ['Class Average (%)', stats.get('class_average', 0)],
            ['Highest (%)', stats.get('highest_score', 0)],
            ['Lowest (%)', stats.get('lowest_score', 0)],
            ['Passing Students', stats.get('passing_students', 0)],
            ['Failing Students', stats.get('failing_students', 0)],
        ]
        stats_table = Table(stats_rows, colWidths=[55*mm, 105*mm])
        stats_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
        ]))
        elements.extend([Paragraph('Statistics', styles['Heading2']), stats_table, Spacer(1, 8)])

        # Student marks table
        header = ['Student', 'Roll', 'Assessment', 'Marks', 'Max', 'Percent']
        rows = [header]
        for sm in report.get('student_marks', []):
            for m in sm.get('marks', []):
                percent = m.get('percentage') if 'percentage' in m else (
                    round((m.get('marks_obtained', 0) / m.get('max_marks', 1)) * 100, 2) if m.get('max_marks') else 0
                )
                rows.append([
                    sm.get('student_name',''), sm.get('roll_number',''), m.get('assessment_type',''),
                    m.get('marks_obtained',''), m.get('max_marks',''), percent
                ])
        if len(rows) == 1:
            rows.append(['No data', '', '', '', '', ''])
        # Wrap text in table data
        rows_wrapped = ReportingService._wrap_table_data(rows, skip_header=True, header_text_white=True)
        page_width = A4[0] - (18*mm + 18*mm)
        tbl = Table(rows_wrapped, repeatRows=1, colWidths=ReportingService._full_width_colwidths(page_width, len(rows[0])))
        tbl.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
            ('BACKGROUND', (0,0), (-1,0), colors.black),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold')
        ]))
        elements.append(tbl)

        doc.build(elements)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    @staticmethod
    def generate_class_attendance_report_pdf(report):
        """Generate a PDF for the class attendance report and return bytes."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
        elements = []
        styles = getSampleStyleSheet()
        # Header (same format as student PDF)
        try:
            from flask import current_app
            logo_path = current_app.root_path + '/static/img/logo-removebg-preview.png'
            logo_img = Image(logo_path)
            logo_img._restrictSize(26*mm, 26*mm)
        except Exception:
            logo_img = ''
        from reportlab.lib.styles import ParagraphStyle
        header_title = ParagraphStyle('HeaderTitle', parent=styles['Title'], alignment=0, fontSize=16, leading=19)
        header_sub = ParagraphStyle('HeaderSub', parent=styles['Normal'], alignment=0, fontSize=10, leading=12)
        header_text = [
            Paragraph('Dr. B. B. Hegde First Grade College, Kundapura', header_title),
            Paragraph('A Unit of Coondapur Education Society (R)', header_sub)
        ]
        header_table = Table([[logo_img, header_text]], colWidths=[26*mm, 148*mm])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LINEBELOW', (0,0), (-1,0), 0.75, colors.lightgrey),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(header_table)
        elements.append(Spacer(1, 6))
        elements.append(Paragraph('Class Attendance Report', styles['Title']))
        s = report.get('subject', {})
        meta_rows = [
            ['Subject', f"{s.get('name','')} ({s.get('code','')})"],
            ['Course', s.get('course_display') or s.get('course') or '']
        ]
        if s.get('section'):
            meta_rows.append(['Section', str(s.get('section')).upper()])
        meta_rows.append(['Year/Sem', f"{s.get('year','')}/{s.get('semester','')}"])
        # Faculty line if available
        try:
            lecturers = s.get('lecturers') or []
            if lecturers:
                meta_rows.append(['Faculty', ', '.join(lecturers)])
        except Exception:
            pass
        # Period after faculty
        meta_rows.append(['Period', f"{report.get('month','')} {report.get('year','')}"])
        meta_table = Table(meta_rows, colWidths=[40*mm, 120*mm])
        meta_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
        ]))
        elements.extend([Spacer(1, 6), meta_table, Spacer(1, 8)])

        stats = report.get('statistics', {})
        stats_rows = [
            ['Total Students', stats.get('total_students', 0)],
            ['Classes Conducted', stats.get('total_classes_conducted', 0)],
            ['Class Average (%)', stats.get('class_average_attendance', 0)],
            ['Good Attendance (≥75%)', stats.get('students_with_good_attendance', 0)],
            ['Poor Attendance (<50%)', stats.get('students_with_poor_attendance', 0)],
        ]
        stats_table = Table(stats_rows, colWidths=[60*mm, 100*mm])
        stats_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
        ]))
        elements.extend([Paragraph('Statistics', styles['Heading2']), stats_table, Spacer(1, 8)])

        header = ['Student', 'Roll', 'Total', 'Present', 'Absent', 'Percent', 'Status']
        rows = [header]
        for st in report.get('student_attendance', []):
            rows.append([
                st.get('student_name',''), st.get('roll_number',''), st.get('total_classes',''),
                st.get('present_classes',''), st.get('absent_classes',''), st.get('attendance_percentage',''), st.get('status','')
            ])
        if len(rows) == 1:
            rows.append(['No data', '', '', '', '', '', ''])
        # Wrap text in table data
        rows_wrapped = ReportingService._wrap_table_data(rows, skip_header=True, header_text_white=True)
        page_width = A4[0] - (18*mm + 18*mm)
        tbl = Table(rows_wrapped, repeatRows=1, colWidths=ReportingService._full_width_colwidths(page_width, len(rows[0])))
        tbl.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
            ('BACKGROUND', (0,0), (-1,0), colors.black),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold')
        ]))
        elements.append(tbl)

        doc.build(elements)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    @staticmethod
    def generate_course_overview_report_pdf(report):
        """Generate a PDF for the course overview report and return bytes."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
        elements = []
        styles = getSampleStyleSheet()
        # Header (same format as student PDF)
        try:
            from flask import current_app
            logo_path = current_app.root_path + '/static/img/logo-removebg-preview.png'
            logo_img = Image(logo_path)
            logo_img._restrictSize(26*mm, 26*mm)
        except Exception:
            logo_img = ''
        from reportlab.lib.styles import ParagraphStyle
        header_title = ParagraphStyle('HeaderTitle', parent=styles['Title'], alignment=0, fontSize=16, leading=19)
        header_sub = ParagraphStyle('HeaderSub', parent=styles['Normal'], alignment=0, fontSize=10, leading=12)
        header_text = [
            Paragraph('Dr. B. B. Hegde First Grade College, Kundapura', header_title),
            Paragraph('A Unit of Coondapur Education Society (R)', header_sub)
        ]
        header_table = Table([[logo_img, header_text]], colWidths=[26*mm, 148*mm])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LINEBELOW', (0,0), (-1,0), 0.75, colors.lightgrey),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(header_table)
        elements.append(Spacer(1, 6))
        elements.append(Paragraph('Course Overview Report', styles['Title']))
        c = report.get('course', {})
        meta_rows = [
            ['Course', f"{c.get('name','')} ({c.get('code','')})"],
            ['Duration (years)', c.get('duration_years','')],
            ['Total Semesters', c.get('total_semesters','')],
            ['Total Students', report.get('total_students','')],
            ['Total Subjects', report.get('total_subjects','')],
        ]
        meta_table = Table(meta_rows, colWidths=[55*mm, 105*mm])
        meta_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
        ]))
        elements.extend([Spacer(1, 6), meta_table, Spacer(1, 8)])

        # Subjects table
        header = ['Subject', 'Code', 'Y/S', 'Enrolled', 'Avg Marks %', 'Pass Rate %', 'Avg Attendance %']
        rows = [header]
        for s in report.get('subjects', []):
            rows.append([
                s.get('subject_name',''), s.get('subject_code',''), f"{s.get('year','')}/{s.get('semester','')}",
                s.get('enrolled_students',''), s.get('marks_statistics',{}).get('average_marks',0),
                s.get('marks_statistics',{}).get('passing_rate',0), s.get('attendance_statistics',{}).get('average_attendance',0)
            ])
        if len(rows) == 1:
            rows.append(['No data', '', '', '', '', '', ''])
        # Wrap text in table data
        rows_wrapped = ReportingService._wrap_table_data(rows)
        page_width = A4[0] - (18*mm + 18*mm)
        tbl = Table(rows_wrapped, repeatRows=1, colWidths=ReportingService._full_width_colwidths(page_width, len(rows[0])))
        tbl.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('INNERGRID', (0,0), (-1,-1), 0.25, colors.lightgrey),
            ('BACKGROUND', (0,0), (-1,0), colors.black),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold')
        ]))
        elements.append(tbl)

        doc.build(elements)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes